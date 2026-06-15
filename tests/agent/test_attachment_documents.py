"""Office 附件文本抽取测试。"""

from __future__ import annotations

import io

from docx import Document
from openpyxl import Workbook

from agentpod_agent.core.attachment_documents import (
    OFFICE_EXTRACTED_MIME,
    extract_document_text,
    has_office_filename,
    is_office_attachment,
    is_office_extracted_blob,
)


def test_is_office_attachment_by_extension() -> None:
    assert is_office_attachment(name="report.xlsx", mime_type=None)
    assert is_office_attachment(name="notes.docx", mime_type=None)
    assert not is_office_attachment(name="readme.txt", mime_type="text/plain")


def test_extract_xlsx_roundtrip() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Amount"])
    ws.append(["Jan", 100])
    ws.append(["Feb", 150])
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    text, truncated, err = extract_document_text(data, name="sales.xlsx")
    assert err is None
    assert truncated is False
    assert "Sales" in text
    assert "Jan" in text
    assert "150" in text


def test_extract_docx_roundtrip() -> None:
    doc = Document()
    doc.add_paragraph("Hello spreadsheet")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    text, truncated, err = extract_document_text(data, name="note.docx")
    assert err is None
    assert truncated is False
    assert "Hello spreadsheet" in text


def test_reject_legacy_xls() -> None:
    _text, _truncated, err = extract_document_text(b"fake", name="old.xls")
    assert err is not None
    assert ".xlsx" in err


def test_is_office_extracted_blob() -> None:
    assert is_office_extracted_blob(name="sales.xlsx", mime_type=OFFICE_EXTRACTED_MIME)
    assert not is_office_extracted_blob(name="readme.txt", mime_type=OFFICE_EXTRACTED_MIME)
    assert not is_office_extracted_blob(name="sales.xlsx", mime_type="application/vnd.ms-excel")


def test_has_office_filename() -> None:
    assert has_office_filename("report.xlsx")
    assert not has_office_filename("notes.txt")
